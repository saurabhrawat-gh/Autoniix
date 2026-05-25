"""
Format file1.xlsx — fixed:
  - Rows 1-3 each fully merged A:F (span linearly), all content preserved
  - Remarks: no wrap, wide column — no overflow/overlap
  - Print: A4 portrait, fit to 1 page
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import WorksheetProperties, PageSetupProperties

INPUT  = "/home/saurabh/Desktop/YouTube/Autonix/file1.xlsx"
OUTPUT = "/home/saurabh/Desktop/YouTube/Autonix/file1_formatted.xlsx"

wb = load_workbook(INPUT)
ws = wb.active

LAST_COL   = 6
LAST_COL_L = "F"

thin  = Side(style="thin")
thick = Side(style="medium")

def bdr(top=None, bottom=None, left=None, right=None):
    return Border(top=top, bottom=bottom, left=left, right=right)

def cell_fmt(cell, size, bold=False, halign="center", wrap=False):
    cell.font      = Font(name="Calibri", size=size, bold=bold)
    cell.alignment = Alignment(horizontal=halign, vertical="center", wrap_text=wrap)

def drop_merges(row):
    """Remove all merges touching this row."""
    to_drop = [str(mc) for mc in ws.merged_cells.ranges
               if mc.min_row <= row <= mc.max_row]
    for ref in to_drop:
        ws.unmerge_cells(ref)

def full_merge(row, value, size, bold=False, halign="center"):
    """Merge A{row}:F{row} and set value + format."""
    drop_merges(row)
    ws.merge_cells(f"A{row}:F{row}")
    c = ws.cell(row=row, column=1)
    c.value = value
    cell_fmt(c, size=size, bold=bold, halign=halign)

# ════════════════════════════════════════════════════════════════════════════
# ROW 1 — title: full merge
# ════════════════════════════════════════════════════════════════════════════
full_merge(1,
    "Working History Sheet of Sh. Iqbal Chand, MCM/WM",
    size=14, bold=True)
ws.row_dimensions[1].height = 24

# ════════════════════════════════════════════════════════════════════════════
# ROW 2 — DOB / DOA / DOS / DOR: single merged banner
# ════════════════════════════════════════════════════════════════════════════
full_merge(2,
    "DOB:- 23.06.1966          DOA:- 12.06.1986          DOS:- 23.12.93          DOR:- 30.06.2026",
    size=11)
ws.row_dimensions[2].height = 18

# ════════════════════════════════════════════════════════════════════════════
# ROW 3 — PF No. / Bill Unit: single merged banner
# ════════════════════════════════════════════════════════════════════════════
full_merge(3,
    "PF No. 50300516363                    Bill Unit= 0301-268",
    size=11, halign="left")
ws.row_dimensions[3].height = 18

# ════════════════════════════════════════════════════════════════════════════
# ROW 4 — thin separator
# ════════════════════════════════════════════════════════════════════════════
ws.row_dimensions[4].height = 4

# ════════════════════════════════════════════════════════════════════════════
# ROW 5 — column headers
# ════════════════════════════════════════════════════════════════════════════
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
col_labels = {
    1: "S. No.",
    2: "Date",
    3: "Pay\n(Substantive)",
    4: "Pay\n(Officiating)",
    5: "Grade/\nLevel",
    6: "Remarks",
}
for col, label in col_labels.items():
    c = ws.cell(row=5, column=col)
    c.value     = label
    c.font      = Font(name="Calibri", size=10, bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.fill      = HEADER_FILL
    c.border    = bdr(top=thick, bottom=thick, left=thin, right=thin)
ws.row_dimensions[5].height = 30

# ════════════════════════════════════════════════════════════════════════════
# ROWS 6+ — data
# ════════════════════════════════════════════════════════════════════════════
max_row = ws.max_row
for row in range(6, max_row + 1):
    for col in range(1, LAST_COL + 1):
        c = ws.cell(row=row, column=col)
        c.font   = Font(name="Calibri", size=9)
        is_remarks = (col == LAST_COL)
        c.alignment = Alignment(
            horizontal="left" if is_remarks else "center",
            vertical="center",
            wrap_text=False      # no wrap → no overlap between rows
        )
        c.border = bdr(top=thin, bottom=thin, left=thin, right=thin)
    ws.row_dimensions[row].height = 13

# ════════════════════════════════════════════════════════════════════════════
# COLUMN WIDTHS
# ════════════════════════════════════════════════════════════════════════════
# Remarks needs to fit longest entry without wrapping
col_widths = {1: 6, 2: 12, 3: 14, 4: 14, 5: 14, 6: 42}
for col, w in col_widths.items():
    ws.column_dimensions[get_column_letter(col)].width = w

# ════════════════════════════════════════════════════════════════════════════
# PRINT SETUP — A4 portrait, fit to 1 page
# ════════════════════════════════════════════════════════════════════════════
ws.sheet_properties = WorksheetProperties(
    pageSetUpPr=PageSetupProperties(fitToPage=True)
)
ws.page_setup.paperSize   = ws.PAPERSIZE_A4
ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
ws.page_setup.fitToHeight = 1
ws.page_setup.fitToWidth  = 1

ws.page_margins = PageMargins(
    left=0.4, right=0.4,
    top=0.5,  bottom=0.5,
    header=0.2, footer=0.2
)

ws.print_area       = f"A1:F{max_row}"
ws.print_title_rows = "1:5"
ws.freeze_panes     = "A6"

wb.save(OUTPUT)
print(f"Saved → {OUTPUT}  (rows 1-{max_row}, cols A-F)")
